import json
from PIL import Image
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from pptx import Presentation
from pptx.util import Inches, Pt as P
from pptx.dml.color import RGBColor as C
from pptx.enum.shapes import MSO_SHAPE
R="../robot_model/04_Renders/"
O="out/"
S=json.load(open(O+"bench.json"));W=json.load(open(O+"ws.json"));SIM=json.load(open(O+"sim.json"))
f=lambda x,n=2:f"{x:.{n}f}"
LB=2.4
rr,rs,pm=S["RRT"],S["RRT*"],S["PRM"]
over=lambda x:int(round((x/LB-1)*100))
SHORT="Motion Planning for a Mobile Manipulator"
TITLE="Motion Planning for a TurtleBot3 + OpenMANIPULATOR-X Mobile Manipulator: A Comparison of RRT, RRT* and PRM"
TEAM=["Kaviyashree T S","Barath K M","Jyothish S","Nabeeha Naaz H"]
SUBJ="Robot Dynamic Control"
REFS=[
"S. M. LaValle, \"Rapidly-exploring random trees: A new tool for path planning,\" Tech. Rep. TR 98-11, Computer Science Dept., Iowa State University, 1998.",
"L. E. Kavraki, P. Svestka, J.-C. Latombe and M. H. Overmars, \"Probabilistic roadmaps for path planning in high-dimensional configuration spaces,\" IEEE Trans. Robotics and Automation, vol. 12, no. 4, pp. 566-580, 1996.",
"S. Karaman and E. Frazzoli, \"Sampling-based algorithms for optimal motion planning,\" Int. J. Robotics Research, vol. 30, no. 7, pp. 846-894, 2011.",
"J. J. Kuffner and S. M. LaValle, \"RRT-Connect: An efficient approach to single-query path planning,\" Proc. IEEE ICRA, pp. 995-1001, 2000.",
"I. A. Sucan, M. Moll and L. E. Kavraki, \"The Open Motion Planning Library,\" IEEE Robotics & Automation Magazine, vol. 19, no. 4, pp. 72-82, 2012.",
"S. Chitta, I. Sucan and S. Cousins, \"MoveIt!,\" IEEE Robotics & Automation Magazine, vol. 19, no. 1, pp. 18-19, 2012.",
"S. Macenski, T. Foote, B. Gerkey, C. Lalancette and W. Woodall, \"Robot Operating System 2: Design, architecture, and uses in the wild,\" Science Robotics, vol. 7, no. 66, 2022.",
"N. Koenig and A. Howard, \"Design and use paradigms for Gazebo, an open-source multi-robot simulator,\" Proc. IEEE/RSJ IROS, pp. 2149-2154, 2004.",
"S. Macenski, F. Martin, R. White and J. G. Clavero, \"The Marathon 2: A Navigation System,\" Proc. IEEE/RSJ IROS, 2020.",
"B. Siciliano, L. Sciavicco, L. Villani and G. Oriolo, Robotics: Modelling, Planning and Control. Springer, 2009.",
"ROBOTIS Co., Ltd., \"TurtleBot3 and OpenMANIPULATOR-X e-Manual,\" https://emanual.robotis.com (accessed 2026).",
"ROBOTIS-GIT, \"turtlebot3_manipulation\" (ROS 2 Humble branch), https://github.com/ROBOTIS-GIT/turtlebot3_manipulation.",
"J. D. Gammell, S. S. Srinivasa and T. D. Barfoot, \"Informed RRT*: Optimal sampling-based path planning focused via direct sampling of an admissible ellipsoidal heuristic,\" Proc. IEEE/RSJ IROS, pp. 2997-3004, 2014.",
]
JT=[("base_fixed","base_link -> link1","(-0.092, 0, 0.091)","-","fixed"),
("joint1","link1 -> link2","(0.012, 0, 0.017)","z","+/-0.9 pi"),
("joint2","link2 -> link3","(0, 0, 0.0595)","y","-0.57 pi ... 0.5 pi"),
("joint3","link3 -> link4","(0.024, 0, 0.128)","y","-0.3 pi ... 0.44 pi"),
("joint4","link4 -> link5","(0.124, 0, 0)","y","-0.57 pi ... 0.65 pi"),
("gripper_left/right","link5 -> finger","(0.0817, +/-0.021, 0)","+/-y","-0.010 ... 0.019 m"),
("end_effector_joint","link5 -> tool","(0.126, 0, 0)","-","fixed"),
("wheel_left/right","base_link -> wheel","(0, +/-0.144, 0.023)","z","continuous")]
LITT=[
("LaValle (1998)\nRapidly-exploring random trees: a new tool for path planning [1]","Introduced RRT: a tree grown by sampling random states and extending the nearest node. Probabilistically complete, works in high-dimensional spaces, very quick to a first path.","Paths are jagged and not optimal; behaviour depends on step size and distance metric."),
("Kavraki et al. (1996)\nProbabilistic roadmaps for path planning in high-dimensional configuration spaces [2]","Introduced PRM: build a roadmap of free configurations once, then answer many queries by graph search. Good for static scenes and many goals.","Roadmap construction is costly; narrow passages are hard; roadmap must be rebuilt if obstacles move."),
("Kuffner & LaValle (2000)\nRRT-Connect: an efficient approach to single-query path planning [4]","Grows two trees (from start and goal) that try to connect greedily; much faster than basic RRT on manipulator-type problems.","Still not optimal, so paths need post-smoothing; only for problems where two states can be connected directly."),
("Karaman & Frazzoli (2011)\nSampling-based algorithms for optimal motion planning [3]","Showed that RRT and PRM are not asymptotically optimal; proposed RRT* and PRM*, which converge to the optimal path with the same asymptotic complexity.","Convergence can be slow in higher dimensions; extra cost per sample from neighbour search and rewiring; radius needs tuning."),
("Gammell et al. (2014)\nInformed RRT* [13]","After the first solution, samples only inside an ellipsoid that can still improve it, so the path converges faster than RRT*.","Needs a first solution; advantage shrinks as dimension grows; assumes a simple distance-based cost."),
("Sucan et al. (2012)\nThe Open Motion Planning Library [5]","Common C++ interface to many planners (RRT, RRT*, PRM and more) with benchmarking tools; it is the planning back-end of MoveIt.","Generic library: results depend on the robot model, collision checker and parameters chosen by the user."),
("Chitta et al. (2012)\nMoveIt! [6]","Brings kinematics, collision checking, planning and execution together for manipulators in ROS.","Needs per-robot configuration; arm-centred, so planning base and arm together takes extra set-up."),
("Macenski et al. (2022)\nRobot Operating System 2 [7]","Describes ROS 2 design: DDS communication, real-time support, multi-robot use and security.","General framework with no manipulation-specific planning comparison."),
("Koenig & Howard (2004)\nDesign and use paradigms for Gazebo [8]","Open-source simulator with physics and sensor models, widely used to test robots before hardware.","Simulation-to-real gap; results are only as good as the URDF, inertia and friction values supplied."),
("ROBOTIS (e-Manual, GitHub) [11], [12]\nturtlebot3_manipulation","Official URDF, MoveIt configuration and Gazebo files for the TurtleBot3 + OpenMANIPULATOR-X.","No editable parametric CAD source and no planner comparison; meshes are fixed."),
]
GAPS=["Many planner studies use fixed-base industrial arms rather than a small ROS 2 mobile manipulator.",
"Papers compare algorithms in general, but give no editable CAD-to-URDF model of the TurtleBot3 + OpenMANIPULATOR-X to repeat the study on.",
"The official ROBOTIS package gives the model and a working MoveIt setup, but no side-by-side RRT / RRT* / PRM comparison and no parametric CAD.",
"RRT, RRT* and PRM are often explained through theory and generic benchmarks; we also show each one running on the same robot with the same metrics."]
NOVEL=["One specification file (robot_spec.py) generates both the SolidWorks model and the URDF, so geometry and kinematics cannot drift apart. Changing one parameter, e.g. forearm length, updates the CAD and the simulation model together.",
"A fully parametric SolidWorks 2025 model of the TurtleBot3 Waffle Pi + OpenMANIPULATOR-X (19 parts, working joint mates, two-finger gripper) with a drop-in ROS 2 package for MoveIt 2 and Gazebo.",
"A like-for-like comparison of RRT, RRT* and PRM on this robot: same kinematics, same collision model, same start/goal and same metrics, with a simulation and a written justification for each algorithm.",
"A reproducible benchmark script that the team can re-run, and a clear path to repeat the same study in MoveIt 2 / OMPL / Gazebo."]
ABSTRACT=("A mobile manipulator is a wheeled robot with an arm on top, and the arm must move around obstacles without touching them. "
"In this project we built a model, ready for simulation, of a TurtleBot3 Waffle Pi carrying an OpenMANIPULATOR-X arm (4 joints plus a two-finger gripper) and used it to study three sampling-based motion-planning algorithms: RRT, RRT* and PRM. "
"The robot was modelled in SolidWorks 2025 as 19 parametric parts and 2 assemblies, exported as a STEP file and converted into a ROS 2 Humble package (URDF/xacro with STL meshes) whose joints follow the official ROBOTIS description. "
"To study the algorithms we wrote the three planners in Python on the arm's kinematic chain and tested them in a scene where the straight path from start to goal is blocked, with 10 random runs per planner. "
f"All three succeeded every time. RRT was the fastest (about {f(rr['time_mean']*1000,0)} ms) but its paths were long ({f(rr['len_mean'])} rad on average). RRT* took about {f(rs['time_mean']*1000,0)} ms for its first path and, given a 10 s budget, improved it to {f(rs['len_mean'])} rad, which is {over(rs['len_mean'])} % above the straight-line lower bound of 2.40 rad. "
f"PRM needed about {f(pm['time_mean'],1)} s to build its roadmap and returned paths of {f(pm['len_mean'])} rad, but the roadmap can be reused for later queries. The full MoveIt 2 / OMPL / Gazebo comparison is the next step.")
BG=[
"A mobile manipulator can drive to a location and then work on objects there, which is useful for fetching, sorting, inspection and service tasks. The TurtleBot3 Waffle Pi is a popular, low-cost ROS 2 platform with a differential drive, an LDS-01 LiDAR, a Raspberry Pi computer, an OpenCR controller and a camera. The OpenMANIPULATOR-X is a 4-DOF arm with a gripper, built from DYNAMIXEL XM430-W210 servos. ROBOTIS offers the two together as the TurtleBot3 Manipulation setup.",
"Moving the arm safely is a motion-planning problem. We need a smooth path in joint space from a start pose to a goal pose that stays clear of obstacles and of the robot's own body. Every extra joint multiplies the number of possible poses, so checking them all is not practical. Sampling-based planners get around this by testing random poses, and they are what MoveIt 2 and the OMPL library use in ROS 2.",
"To test planners in simulation we first need a trustworthy 3-D model of the robot. The choice of planner (RRT, RRT* or PRM) also matters: it changes how long planning takes, how good the path is, and whether the work can be reused. So the project has three parts: (i) build an editable parametric model that matches the real robot's kinematics, (ii) package it for ROS 2, MoveIt 2 and Gazebo, and (iii) study and compare the three planning algorithms.",
]
OBJ=["Build a parametric SolidWorks 2025 model of the TurtleBot3 Waffle Pi with the OpenMANIPULATOR-X, including working revolute joints and a two-finger gripper.",
"Export it as a STEP file and as a ROS 2 Humble description package (URDF, xacro, STL meshes, RViz launch).",
"Check that the kinematics, joint limits and inertial values match the official ROBOTIS model.",
"Explain RRT, RRT* and PRM mathematically, simulate each one on the robot and justify its behaviour.",
"Compare the three on planning time, path length, success rate and node count.",
"Get the package ready for the MoveIt 2 + OMPL + Gazebo comparison."]
ALG={
"RRT":["Start the tree T with q_s.","Sample q_rand in the joint limits (with probability 0.1 take q_g instead).","Find the nearest node q_near in T and step towards q_rand to get q_new (Eq. 6).","If the segment q_near -> q_new is collision-free, add q_new to T.","If q_new can reach q_g, stop and read the path back through the parents."],
"RRT*":["Steps 1-3 as in RRT.","Collect the neighbours Q_near within radius r_n (Eq. 9).","Choose the parent that gives the lowest cost-to-come (Eq. 7).","Rewire: for each neighbour, if going through q_new is cheaper, make q_new its parent (Eq. 8).","Do not stop at the first solution; keep sampling until the time budget ends and keep the cheapest path."],
"PRM":["Learning phase: sample N collision-free configurations (plus q_s and q_g).","Connect each node to its k nearest neighbours if the straight segment is free (Eq. 11).","Query phase: run Dijkstra from q_s to q_g over the roadmap (Eq. 12).","If no path exists, add more nodes and repeat."]}
METH_CAD=[
("4.2 CAD modelling","We modelled the robot in SolidWorks 2025, driving the SolidWorks API from Python scripts (tools/build_parts.py and build_robot.py). Every part is parametric: its dimensions are tied to Global Variables (Tools > Equations), for example J3_Z (upper-arm length), J4_X (forearm length), Hole_P (waffle-plate hole pitch) and Tire_D (wheel diameter). The model has 19 parts and 2 assemblies. ASM01 is the Waffle Pi base and ASM00 is the complete robot, with its origin on the ground (base_footprint)."),
("4.3 Assembly and joints","Joints 1 to 4 are mated as real revolute joints (a collinear axis mate plus a coincident plane), and the wheels spin on their axle. The two gripper fingers are slider mates with limits of -10 mm to +19 mm around the 21 mm home position, which gives a jaw opening of roughly 6 mm to 64 mm. Parts are modelled in the ROS convention (x forward, y left, z up), so each link's origin sits exactly on its URDF frame."),
("4.4 ROS 2 description","tools/gen_urdf.py writes the URDF from the same specification file (robot_spec.py) that drives the CAD, so the two cannot drift apart. Meshes are exported as binary STL in metres. The inertial values are the official ROBOTIS ones. The base and wheels use the official collision primitives, and the arm links use their meshes. Link and joint names are unchanged, so the official SRDF, ompl_planning.yaml, controllers and Nav2 configuration work without edits."),
("4.5 Verification","The joints (origin, rpy, type and axis) were checked automatically against turtlebot3_manipulation_description from the ROS 2 Humble branch, and the project README records that all of them matched. The package opens in RViz with joint sliders through display.launch.py.")]
# ------------------------------------------------------------- DOCX helpers
d=Document()
sec=d.sections[0];sec.left_margin=sec.right_margin=Cm(2.3);sec.top_margin=sec.bottom_margin=Cm(2.3)
st=d.styles['Normal'];st.font.name='Times New Roman';st.font.size=Pt(12);st.paragraph_format.line_spacing=1.3;st.paragraph_format.space_after=Pt(6)
for n,s in (('Heading 1',16),('Heading 2',13),('Heading 3',12)):
    h=d.styles[n];h.font.name='Times New Roman';h.font.size=Pt(s);h.font.bold=True;h.font.color.rgb=RGBColor(0x1F,0x2A,0x44)
def para(t,al=WD_ALIGN_PARAGRAPH.JUSTIFY,b=False,i=False,sz=None):
    p=d.add_paragraph();p.alignment=al;r=p.add_run(t);r.bold=b;r.italic=i
    if sz:r.font.size=Pt(sz)
    return p
def shade(cell,hexc):
    tcPr=cell._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:val'),'clear');sh.set(qn('w:color'),'auto');sh.set(qn('w:fill'),hexc);tcPr.append(sh)
def table(rows,hdr,widths=None,sz=10,style='Table Grid'):
    t=d.add_table(rows=1,cols=len(hdr));t.style=style;t.autofit=False
    for c,h in zip(t.rows[0].cells,hdr):
        c.text='';r=c.paragraphs[0].add_run(h);r.bold=True;r.font.size=Pt(sz);r.font.color.rgb=RGBColor(255,255,255);shade(c,"1F2A44")
    for k,row in enumerate(rows):
        cs=t.add_row().cells
        for c,v in zip(cs,row):
            c.text=''
            parts=str(v).split("\n")
            for pi,pt in enumerate(parts):
                p=c.paragraphs[0] if pi==0 else c.add_paragraph()
                p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=1.1
                r=p.add_run(pt);r.font.size=Pt(sz)
                if pi==0 and len(parts)>1:r.bold=True
            if k%2==1:shade(c,"F3F5F8")
    if widths:
        for col,w in zip(t.columns,widths):col.width=Cm(w)
        for row in t.rows:
            for c,w in zip(row.cells,widths):c.width=Cm(w)
    d.add_paragraph()
    return t
def fig(path,w,cap):
    d.add_picture(path,width=Cm(w));d.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER
    para(cap,WD_ALIGN_PARAGRAPH.CENTER,i=True,sz=10)
EQN=[0]
def numbered(items):
    for i,t in enumerate(items,1):
        p=d.add_paragraph(f'{i}.  {t}');p.paragraph_format.left_indent=Cm(1);p.paragraph_format.first_line_indent=Cm(-0.7);p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
def eq(name):
    EQN[0]+=1
    im=Image.open(O+f"eq_{name}.png");w=min(14.0,im.width/220*2.54*0.85)
    t=d.add_table(rows=1,cols=2);t.autofit=False
    c0,c1=t.rows[0].cells;t.columns[0].width=Cm(14.2);t.columns[1].width=Cm(1.8);c0.width=Cm(14.2);c1.width=Cm(1.8)
    c0.paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.CENTER;c0.paragraphs[0].add_run().add_picture(O+f"eq_{name}.png",width=Cm(w))
    c1.paragraphs[0].add_run(f"({EQN[0]})");c1.paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.RIGHT
    d.add_paragraph().paragraph_format.space_after=Pt(0)
# ------------------------------------------------------------- DOCX
for _ in range(4):d.add_paragraph()
para("PROJECT REVIEW REPORT",WD_ALIGN_PARAGRAPH.CENTER,b=True,sz=14)
para(TITLE,WD_ALIGN_PARAGRAPH.CENTER,b=True,sz=20)
d.add_paragraph()
d.add_picture(R+"ASM00_iso.png",width=Cm(10));d.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER
d.add_paragraph()
para("Subject: "+SUBJ,WD_ALIGN_PARAGRAPH.CENTER,b=True)
para("Team members",WD_ALIGN_PARAGRAPH.CENTER)
for t in TEAM:para(t,WD_ALIGN_PARAGRAPH.CENTER)
para("Project review: 14 October 2026",WD_ALIGN_PARAGRAPH.CENTER)
d.add_page_break()
d.add_heading("1. Abstract",1);para(ABSTRACT)
para("Keywords: mobile manipulator, TurtleBot3, OpenMANIPULATOR-X, motion planning, RRT, RRT*, PRM, MoveIt 2, OMPL, ROS 2, SolidWorks.",i=True)
d.add_heading("2. Background",1)
for t in BG:para(t)
d.add_heading("2.1 Objectives",2)
numbered(OBJ)
d.add_heading("3. Literature Review",1)
para("Table 1 lists the work this project builds on. For each source we note what it contributed (key findings) and where it falls short for our purpose (limitations / research gap).")
table([(str(i+1),a,b,c) for i,(a,b,c) in enumerate(LITT)],["No.","Title","Key findings","Limitations / Research gap"],[0.9,4.6,5.5,5.2],sz=9)
para("Table 1. Literature summary.",WD_ALIGN_PARAGRAPH.CENTER,i=True,sz=10)
d.add_heading("3.1 Research gap",2)
for t in GAPS:d.add_paragraph(t,style='List Bullet')
d.add_heading("3.2 Novelty of this work",2)
para("To the best of our knowledge, and based on the sources above, the contributions of this project are:")
numbered(NOVEL)
d.add_heading("4. Methodology",1)
d.add_heading("4.1 System architecture",2)
para("The project has four layers (Figure 1). The design layer holds the single specification file (robot_spec.py) and the SolidWorks model, and produces the STEP file and renders. The description layer turns the same specification into a URDF/xacro package with STL meshes that opens in RViz. These two layers are finished. The planning/control layer (MoveIt 2 with OMPL planners and ros2_control controllers) and the simulation/hardware layer (Gazebo, Nav2, the real robot) are the next step, and the package is already prepared for them: link and joint names match the official ROBOTIS files, so the official SRDF, controllers and ompl_planning.yaml work unchanged.")
fig(O+"arch_system.png",15.5,"Figure 1. System architecture. Solid boxes are done, dashed boxes are planned.")
para("The algorithm study in this report runs on the pipeline in Figure 2. The planner proposes joint-space configurations. The collision checker turns each one into link positions with forward kinematics and tests them against the obstacles. The planner then returns a path, and we record time, path length, success and node count.")
fig(O+"arch_pipeline.png",15.5,"Figure 2. Planning pipeline used in the benchmark.")
for h,t in METH_CAD:
    d.add_heading(h,2);para(t)
    if h.startswith("4.3"):fig(R+"ASM00_gripper_iso.png",8.5,"Figure 3. Two-finger gripper on link 5 (SolidWorks render).")
    if h.startswith("4.4"):
        table(JT,["Joint","Parent -> child","Origin (m)","Axis","Limits"],[3.4,3.6,3.8,1.3,3.8])
        para("Table 2. Kinematic chain, identical to the official ROBOTIS model.",WD_ALIGN_PARAGRAPH.CENTER,i=True,sz=10)
d.add_heading("4.6 Mathematical formulation",2)
d.add_heading("Forward kinematics",3)
para("Each arm joint i is placed at offset o_i (Table 2) from its parent and rotates about its axis a_i: the z axis for joint 1 and the y axis for joints 2 to 4. The end-effector pose is the product of these transforms, with the rotations given by Eq. (2).")
eq("fk");eq("rot")
d.add_heading("Configuration space and collision",3)
para("The planner works in the 4-D joint space Q. A configuration q is free if it respects the joint limits, keeps the upper-arm, forearm and wrist/tool segments (each treated as a capsule of radius r_l = 25 mm) away from every obstacle sphere (centre c_j, radius r_j), and keeps every point from joint 2 onward at least z_min = 20 mm above the base_link plane. The two lowest links and the gripper fingers are not checked in this simplified model. The distance from a link segment to a sphere centre is the point-to-segment distance of Eq. (4). The cost of a path is its joint-space length, Eq. (5).")
eq("cfree");eq("dist");eq("len")
d.add_heading("RRT",3)
para("RRT extends the tree by at most a step eta towards a random sample, Eq. (6). Here q_rand is sampled uniformly inside the joint limits, and with probability 0.1 it is replaced by q_g (goal bias). RRT is probabilistically complete: the probability of finding a path if one exists tends to 1 as the number of samples grows, but it gives no guarantee about path quality.")
eq("steer")
d.add_heading("RRT*",3)
para("RRT* keeps a cost-to-come c(q) for every node. A new node chooses the parent that minimises cost among its near neighbours, Eq. (7), and then neighbours are rewired through the new node when that lowers their cost, Eq. (8). The neighbour radius shrinks with the number of nodes n, Eq. (9); in our code we use gamma = 2(1 + 1/d)^(1/d) with d = 4, leave out the volume factor, and cap the radius at 1.5 rad (a practical simplification). With this rewiring the path cost converges almost surely to the optimum c*, Eq. (10) [3].")
eq("parent");eq("rewire");eq("radius");eq("opt")
d.add_heading("PRM",3)
para("PRM builds a graph whose vertices are free configurations and whose edges connect each node to its k nearest neighbours when the straight segment between them is free, Eq. (11). A query is then a shortest-path problem solved by Dijkstra's algorithm, Eq. (12). Like RRT it is probabilistically complete, but the path can only follow roadmap edges, so its quality depends on how dense the roadmap is.")
eq("prm");eq("dij")
d.add_heading("Dynamics (how a planned path is executed)",3)
para("Since this subject is about robot dynamics and control, here is how the planner connects to it. The planner only gives a geometric path. In the full stack that path is given a timing and followed by the arm controller in ros2_control, and the arm's motion obeys the manipulator equation (13) with inertia matrix M, Coriolis/centrifugal matrix C and gravity vector g. A computed-torque controller, Eq. (14), cancels these terms and makes the tracking error e decay with gains Kp and Kd. We do not simulate the dynamics in this report, but the inertial values in the URDF are the official ROBOTIS ones, so the model is ready for it.")
eq("dyn");eq("ctc")
d.add_heading("4.7 The three algorithms step by step",2)
for k,v in ALG.items():
    d.add_heading(k,3)
    numbered(v)
d.add_heading("4.8 Benchmark setup",2)
para("We implemented forward kinematics, the capsule collision check and the three planners in Python (bench.py). The obstacles are three spheres in front of the base: one of radius 50 mm at (0.09, 0, 0.21) m and two of radius 40 mm at (0.10, +/-0.12, 0.14) m. The start pose is q_s = (-1.2, -0.2, 0.6, 0.3) rad and the goal is q_g = (1.2, -0.2, 0.6, 0.3) rad, so joint 1 must swing 2.4 rad. We checked that the straight joint-space line from q_s to q_g is in collision, so the planners have to find a detour; the straight-line length of 2.40 rad is therefore a lower bound for any path. Joint limits and the minimum height are enforced, and each edge is collision-checked every 0.05 rad. RRT uses eta = 0.4 rad, 10 % goal bias and stops at its first solution. RRT* uses the same eta with rewiring and a 10 s time budget. PRM uses 1500 nodes, k = 10 and a Dijkstra query. Each planner ran 10 times with different random seeds. We recorded success rate, time to first solution (for PRM: roadmap construction plus query), joint-space path length and node count.")
d.add_heading("4.9 Planned MoveIt 2 / Gazebo study",2)
para("The next step is to repeat the experiment in the full stack with ros2 launch turtlebot3_manipulation_moveit_config moveit_gazebo.launch.py, choosing RRTkConfigDefault, RRTstarkConfigDefault and PRMkConfigDefault in the RViz MotionPlanning panel. That replaces our simplified collision model with the real mesh geometry and Gazebo physics.")
d.add_heading("5. Results",1)
d.add_heading("5.1 CAD model and ROS 2 package",2)
para("We delivered 19 parametric parts (13 TurtleBot3 parts, TB01-TB13, and 6 OpenMANIPULATOR-X parts, OM01-OM06), two SolidWorks assemblies, one STEP file of the whole robot, and a ROS 2 package with 19 STL meshes. The package contains a standalone URDF (20 links and 19 joints, one of which is the mimic joint for the right gripper finger) and a drop-in xacro that reuses the official Gazebo and ros2_control macros. Figure 4 shows the assembly.")
fig(R+"ASM00_iso.png",9,"Figure 4. Complete mobile manipulator (isometric).")
t=d.add_table(rows=1,cols=3)
for c,im in zip(t.rows[0].cells,("ASM00_front.png","ASM00_side.png","ASM00_top.png")):
    c.paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.CENTER;c.paragraphs[0].add_run().add_picture(R+im,width=Cm(5.2))
para("Figure 5. Front, side and top views.",WD_ALIGN_PARAGRAPH.CENTER,i=True,sz=10)
d.add_heading("5.2 Arm workspace",2)
para(f"Sampling 20 000 random configurations within the joint limits and running forward kinematics gives a maximum end-effector distance of {f(W['reach_max']*100,1)} cm from the arm's base frame (link1), a maximum height of {f(W['zmax']*100+1,1)} cm above the floor (base_link is 1 cm above the floor, so this is {f(W['zmax']*100,1)} cm above base_link) and a maximum forward reach of {f(W['xmax']*100,1)} cm in front of the base_link origin (Figure 6). The arm base sits 9.2 cm behind the base_link origin, so the gripper can reach past the front edge of the 266 mm waffle plate.")
fig(O+"workspace.png",14,"Figure 6. Sampled end-effector workspace.")
d.add_heading("5.3 Simulation of each algorithm",2)
para("Note: the planners below are our own Python implementations of the three algorithms, not OMPL, so absolute times depend on our code and the relative trends are what matters. Each figure shows (a) the search structure projected on the q1-q2 plane over a map of colliding configurations (grey, with q3 and q4 fixed at 0.6 and 0.3 rad), (b) arm poses along the returned path in side view, and (c) the four joint angles along the path. In (a) the black dot is the start, the star is the goal and the black line is the path; the path also moves q3 and q4, so it can look as if it crosses the grey area.",i=True)
sr=SIM
def simblock(name,fn,cap,txt):
    d.add_heading(name,3);fig(O+fn,16,cap)
    for t in txt:para(t)
simblock("5.3.1 RRT","sim_rrt.png","Figure 7. RRT simulation (one run).",[
f"RRT reached the goal after growing a tree of only {sr['rrt_nodes']} nodes in this run, and returned a path with {sr['wps']['RRT']} waypoints. Over 10 runs it needed {f(rr['time_mean']*1000,0)} +/- {f(rr['time_sd']*1000,0)} ms and about {rr['nodes_mean']:.0f} nodes.",
f"Justification. RRT always grows from the nearest tree node, which pulls the tree quickly into unexplored space, so it finds a way round the blocking sphere with very few samples. It stops as soon as it reaches the goal and never improves the path, so the route is whatever the random samples happened to give: {f(rr['len_mean'])} +/- {f(rr['len_sd'])} rad, which is {over(rr['len_mean'])} % longer than the 2.40 rad lower bound and varies a lot between runs. The jagged joint curves in panel (c) show the same effect. That is the speed-versus-quality trade-off: RRT is good for a quick first answer and poor as a final one."])
simblock("5.3.2 RRT*","sim_rrtstar.png","Figure 8. RRT* simulation after a 10 s budget (one run).",[
f"RRT* found its first path after about {f(rs['time_mean']*1000,0)} ms, almost as quickly as RRT, and then kept improving it. After 10 s the tree held about {rs['nodes_mean']:.0f} nodes and the path was {f(rs['len_mean'])} +/- {f(rs['len_sd'])} rad, only {over(rs['len_mean'])} % above the lower bound. The path in this run has {sr['wps']['RRT*']} waypoints and smooth joint curves.",
"Justification. Picking the best parent (Eq. 7) and rewiring neighbours (Eq. 8) keeps cutting out detours, so the cost falls over time (Figure 9) and the result is shorter and more repeatable than RRT. The price is computation: each new node looks at all its neighbours, so RRT* builds a tree about 46 times larger than RRT and uses its whole time budget. If time is tight it can be stopped early and still returns a valid path."])
fig(O+"rrtstar_convergence.png",10.5,"Figure 9. RRT* path length against planning time for five runs. Each curve starts when a first path is found.")
simblock("5.3.3 PRM","sim_prm.png","Figure 10. PRM simulation (one run, 1500-node roadmap).",[
f"PRM sampled {sr['rrt_nodes']*0+1500} collision-free nodes and joined them with {sr['prm_edges']} edges in this run; Dijkstra then returned a path with {sr['wps']['PRM']} waypoints. Over 10 runs it took {f(pm['time_mean'],1)} +/- {f(pm['time_sd'],1)} s, almost all of it roadmap construction, and returned paths of {f(pm['len_mean'])} +/- {f(pm['len_sd'])} rad.",
"Justification. PRM spends its time up front testing thousands of edges, so the first query is slow, but any later start-goal pair only needs a graph search. The path can only follow edges between stored nodes, so it is shorter than RRT's but cannot match RRT*'s, which keeps refining. Table 3 shows the effect of roadmap size."])
prm=sr["prm"]
table([(str(x["N"]),f"{x['success']}/10",f(x["time"],1),f(x["len"])) for x in prm],["Nodes N","Success","Build + query time (s)","Mean path length (rad)"],[3,3,4.5,4.5])
para("Table 3. PRM roadmap size study (10 runs per size; 1500-node row from the main benchmark). Build time rises steadily with N. Path length tends to fall as N grows (about 4.1-4.2 rad at N = 100-250 against 3.5-3.6 rad at N = 1000-1500), but neighbouring sizes differ by less than the run-to-run spread (about 0.27 rad).",WD_ALIGN_PARAGRAPH.CENTER,i=True,sz=10)
d.add_heading("5.4 Comparison",2)
rows=[]
for n_ in ("RRT","RRT*","PRM"):
    s=S[n_];rows.append((n_,f"{s['success']}/{s['runs']}",f"{f(s['time_mean']*1000,0)} +/- {f(s['time_sd']*1000,0)}",f"{f(s['len_mean'])} +/- {f(s['len_sd'])}",f"{over(s['len_mean'])} %",f"{s['nodes_mean']:.0f}"))
table(rows,["Algorithm","Success","Time to first solution (ms)","Path length (rad)","Above lower bound","Nodes"],[2.4,2,3.6,3.2,2.6,2])
para("Table 4. Results over 10 runs (mean +/- standard deviation). The lower bound is the 2.40 rad straight line, which is blocked. RRT* length is the best path after its 10 s budget; PRM time is roadmap construction plus query.",WD_ALIGN_PARAGRAPH.CENTER,i=True,sz=10)
fig(O+"bench.png",15,"Figure 11. Algorithm comparison.")
fig(O+"paths.png",10,"Figure 12. Example arm motions planned by each algorithm around the three obstacles.")
d.add_heading("5.5 Discussion and justification",2)
for t in [f"All three algorithms solved the problem in every run, as their completeness properties predict. They differ in what it costs to get there: RRT was quickest ({f(rr['time_mean']*1000,0)} ms) but gave the longest, least consistent paths; RRT* gave the best paths ({f(rs['len_mean'])} rad, {f((1-rs['len_mean']/rr['len_mean'])*100,0)} % shorter than RRT) at the price of using its whole time budget; PRM sat in between on quality ({f(pm['len_mean'])} rad) and was by far the slowest for a single query because it first builds the roadmap.",
"The right algorithm depends on the task. A mobile manipulator keeps arriving at new places, so the scene keeps changing and a one-off planner (RRT or RRT*) fits best; pick RRT* when path quality or smooth joint motion matters. For repeated moves between known poses in a fixed scene, such as a workstation, PRM pays back its set-up cost. So we expect RRT* to be the best default for this robot, and the MoveIt 2 study will check that.",
"Limitations. Our scene has only three spheres, and the planners are plain Python rather than OMPL, so timings are not comparable with MoveIt numbers. The base is held fixed, and we did not model self-collision, the gripper fingers or the real mesh geometry. In our RRT* the costs of a node's descendants are not updated after a rewire, so its internal cost values are approximate; the reported path lengths are measured on the final path itself and are not affected. Ten runs per planner give only a rough picture of the spread. These limits are why we plan the MoveIt 2 / Gazebo study."]:para(t)
d.add_heading("5.6 Conclusion and future work",2)
para("We produced a parametric CAD and ROS 2 model of the TurtleBot3 + OpenMANIPULATOR-X mobile manipulator whose kinematics match the official one, and we simulated and justified three planning algorithms on it. The results show the expected trade-off: RRT is quickest, RRT* gives the best and most consistent paths, and PRM pays up front so it can be reused. Future work: (1) run RRTkConfigDefault, RRTstarkConfigDefault and PRMkConfigDefault in MoveIt 2 with Gazebo on cluttered tables and shelves; (2) add RRT-Connect and Informed RRT*; (3) simulate the arm dynamics and a computed-torque controller tracking the planned path; (4) plan base and arm together; (5) test pick-and-place on the physical robot.")
d.add_heading("6. References",1)
for i,r in enumerate(REFS,1):
    p=d.add_paragraph(f"[{i}] {r}");p.paragraph_format.left_indent=Cm(1);p.paragraph_format.first_line_indent=Cm(-1);p.alignment=WD_ALIGN_PARAGRAPH.LEFT
d.save(O+"Project_Review_Report.docx")
json.dump(dict(neqs=EQN[0]),open(O+"neq.json","w"))
# ------------------------------------------------------------- PPTX
NAVY=C(0x1F,0x2A,0x44);TEAL=C(0x1B,0x9E,0x77);GREY=C(0x55,0x5B,0x66);LIGHT=C(0xF3,0xF5,0xF8);INK=C(0x22,0x26,0x2E)
pr=Presentation();pr.slide_width=Inches(13.333);pr.slide_height=Inches(7.5);blank=pr.slide_layouts[6];n=[0]
def tb(s,x,y,w,h,text,size=18,bold=False,color=NAVY,al=None):
    b=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=b.text_frame;tf.word_wrap=True
    for k,l in enumerate(text if isinstance(text,list) else [text]):
        p=tf.paragraphs[0] if k==0 else tf.add_paragraph();r=p.add_run();r.text=l;r.font.size=P(size);r.font.bold=bold;r.font.color.rgb=color;r.font.name="Calibri";p.space_after=P(8)
        if al:p.alignment=al
    return b
def slide(title):
    s=pr.slides.add_slide(blank);n[0]+=1
    bar=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,Inches(0.25),Inches(7.5));bar.fill.solid();bar.fill.fore_color.rgb=TEAL;bar.line.fill.background()
    tb(s,0.7,0.3,12,0.9,title,30,True)
    tb(s,0.7,7.05,9,0.4,SHORT+"  |  Project review, 14.10.2026",11,color=GREY);tb(s,12.2,7.05,0.8,0.4,str(n[0]),11,color=GREY)
    return s
def bullets(s,items,x=0.7,y=1.4,w=12,size=20,gap=10):
    b=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(5));tf=b.text_frame;tf.word_wrap=True
    for k,t in enumerate(items):
        p=tf.paragraphs[0] if k==0 else tf.add_paragraph();r=p.add_run();r.text="•  "+t;r.font.size=P(size);r.font.color.rgb=INK;r.font.name="Calibri";p.space_after=P(gap)
def tbl(s,rows,hdr,x,y,w,colw,size=13,rh=0.4):
    t=s.shapes.add_table(len(rows)+1,len(hdr),Inches(x),Inches(y),Inches(w),Inches(rh*(len(rows)+1))).table
    for j,h in enumerate(hdr):
        c=t.cell(0,j);c.text=h;c.fill.solid();c.fill.fore_color.rgb=NAVY
        for p in c.text_frame.paragraphs:
            for r in p.runs:r.font.size=P(size);r.font.bold=True;r.font.color.rgb=C(255,255,255)
    for i,row in enumerate(rows,1):
        for j,v in enumerate(row):
            c=t.cell(i,j);c.text="";c.fill.solid();c.fill.fore_color.rgb=LIGHT if i%2 else C(255,255,255)
            parts=str(v).split("\n")
            for pi,pt in enumerate(parts):
                p=c.text_frame.paragraphs[0] if pi==0 else c.text_frame.add_paragraph();r=p.add_run();r.text=pt;r.font.size=P(size);r.font.color.rgb=INK;r.font.bold=(pi==0 and len(parts)>1)
            c.margin_top=c.margin_bottom=Inches(0.03)
    for j,cw in enumerate(colw):t.columns[j].width=Inches(cw)
def pic(s,path,x,y,w=None,h=None):s.shapes.add_picture(path,Inches(x),Inches(y),Inches(w) if w else None,Inches(h) if h else None)
def eqpic(s,name,x,y,maxw):
    im=Image.open(O+f"eq_{name}.png");w=min(maxw,im.width/220*1.25);pic(s,O+f"eq_{name}.png",x,y,w=w);return im.height/im.width*w
# title
s=pr.slides.add_slide(blank);n[0]+=1
bg=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,pr.slide_width,pr.slide_height);bg.fill.solid();bg.fill.fore_color.rgb=NAVY;bg.line.fill.background()
tb(s,0.8,1.0,7.2,3,"Motion Planning for a TurtleBot3 + OpenMANIPULATOR-X Mobile Manipulator",36,True,C(255,255,255))
tb(s,0.8,4.0,7.2,0.8,"Comparison of RRT, RRT* and PRM",22,False,C(0x8F,0xE3,0xC8))
tb(s,0.8,5.0,7.2,2,["Subject: "+SUBJ,"Team: "+", ".join(TEAM),"Project review  |  14 October 2026"],15,False,C(0xD0,0xD5,0xE0))
pic(s,R+"ASM00_iso.png",7.6,1.2,w=5.3)
s=slide("Outline");bullets(s,["1  Abstract","2  Background and objectives","3  Literature review, research gap and novelty","4  Methodology: architecture, equations, algorithms","5  Results: simulation of each algorithm and justification","6  Conclusion and future work","7  References"],size=24)
s=slide("1. Abstract")
bullets(s,["Goal: study three motion-planning algorithms (RRT, RRT*, PRM) on a mobile manipulator",
"Built a parametric SolidWorks 2025 model (19 parts, 2 assemblies, gripper) and a ROS 2 Humble package from one spec file",
"Wrote and simulated each algorithm on the arm; scene where the straight path is blocked; 10 runs each; all succeeded",
f"RRT: fastest ({f(rr['time_mean']*1000,0)} ms) but longest paths ({f(rr['len_mean'])} rad)",
f"RRT*: best paths ({f(rs['len_mean'])} rad, {over(rs['len_mean'])} % above the 2.40 rad lower bound) with a 10 s budget",
f"PRM: {f(pm['time_mean'],1)} s to build the roadmap, {f(pm['len_mean'])} rad paths, reusable",
"Next: full MoveIt 2 + OMPL + Gazebo study"],size=19,gap=8)
s=slide("2. Background")
bullets(s,["Mobile manipulation: drive to a place, then work on objects there","TurtleBot3 Waffle Pi: differential drive, LDS-01 LiDAR, Raspberry Pi 4, OpenCR, camera","OpenMANIPULATOR-X: 4-DOF arm + gripper, DYNAMIXEL XM430-W210 servos","The arm needs collision-free paths in joint space, and the search space grows with every joint","Sampling-based planners in OMPL / MoveIt 2 are the usual answer"],w=7.3,size=19)
pic(s,R+"ASM00_iso.png",8.0,1.6,w=4.9)
s=slide("2. Problem statement and objectives");bullets(s,OBJ,size=19)
for k,(lo,hi) in enumerate(((0,4),(4,7),(7,10))):
    s=slide(f"3. Literature review ({k+1}/3)")
    tbl(s,[(a_,b_,c_) for a_,b_,c_ in LITT[lo:hi]],["Title","Key findings","Limitations / Research gap"],0.5,1.2,12.3,[3.5,4.6,4.2],size=12,rh=1.0)
s=slide("3. Research gap");bullets(s,GAPS,size=21,gap=16)
s=slide("3. Novelty of this work");bullets(s,NOVEL,size=19,gap=14);tb(s,0.7,6.5,12,0.4,"To the best of our knowledge, based on the sources reviewed.",12,False,GREY)
s=slide("4. System architecture");pic(s,O+"arch_system.png",1.4,1.15,h=5.8)
s=slide("4. Planning pipeline used in the benchmark");pic(s,O+"arch_pipeline.png",0.8,1.7,w=11.7)
bullets(s,["Collision checker: forward kinematics -> link capsules (r = 25 mm) vs obstacle spheres","Metrics: success, time to first solution, joint-space path length, node count"],y=5.2,size=18)
s=slide("4. Methodology: CAD and ROS 2 model")
bullets(s,["19 parametric parts: TB01-TB13 (TurtleBot3), OM01-OM06 (arm)","Dimensions tied to Global Variables, e.g. J3_Z, J4_X, Hole_P, Tire_D","Joints 1-4: real revolute mates; fingers: slider mates (6-64 mm opening)","URDF generated from the same spec file, so CAD and URDF stay in step","Same link / joint names as the official package"],w=6.6,size=18)
pic(s,R+"ASM00_gripper_iso.png",7.5,1.6,w=5.4)
s=slide("4. Kinematic chain");tbl(s,JT,["Joint","Parent -> child","Origin (m)","Axis","Limits"],0.7,1.4,11.9,[2.6,2.6,2.8,0.9,3.0],size=13,rh=0.45)
tb(s,0.7,6.2,12,0.6,"Same as turtlebot3_manipulation_description (ROS 2 Humble); checked automatically.",15,False,GREY)
def eqslide(title,items,note=None):
    s=slide(title);y=1.3
    for nm,lab in items:
        tb(s,0.7,y,12,0.35,lab,13,True,TEAL);h=eqpic(s,nm,0.9,y+0.32,11.5);y+=h+0.55
    if note:tb(s,0.7,6.55,12,0.4,note,12,False,GREY)
eqslide("4. Equations (1/4): kinematics and free space",[("fk","Forward kinematics"),("rot","Joint rotations"),("cfree","Free configuration space")])
eqslide("4. Equations (2/4): collision, cost and RRT",[("dist","Point-to-segment distance (collision test)"),("len","Path length (cost)"),("steer","RRT: extend towards a random sample"),("parent","RRT*: choose the cheapest parent")])
eqslide("4. Equations (3/4): RRT* and PRM",[("rewire","RRT*: rewire neighbours"),("radius","RRT*: neighbour radius"),("opt","RRT*: asymptotic optimality"),("prm","PRM: roadmap edges")])
eqslide("4. Equations (4/4): PRM query and dynamics",[("dij","PRM: Dijkstra query"),("dyn","Manipulator dynamics (tracking the planned path)"),("ctc","Computed-torque control")],"Dynamics are shown for context; they are not simulated in this review.")
s=slide("4. The three algorithms")
for k,(nm,steps) in enumerate(ALG.items()):
    x=0.5+k*4.2
    b=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(1.3),Inches(4.0),Inches(5.4));b.fill.solid();b.fill.fore_color.rgb=LIGHT;b.line.color.rgb=TEAL
    tb(s,x+0.15,1.35,3.7,0.5,nm,22,True,TEAL)
    bullets(s,[f"{i+1}. {t}" for i,t in enumerate(steps)],x=x+0.1,y=1.95,w=3.8,size=14,gap=6)
s=slide("4. Benchmark setup")
bullets(s,["Python implementation of FK, capsule collision check and the three planners",
"3 spherical obstacles; straight line start-to-goal is blocked (checked), so a detour is required",
"Start (-1.2, -0.2, 0.6, 0.3) to goal (1.2, -0.2, 0.6, 0.3) rad; lower bound 2.40 rad",
"RRT: step 0.4 rad, 10 % goal bias, stop at first solution",
"RRT*: same step, rewiring, 10 s budget",
"PRM: 1500 nodes, k = 10, Dijkstra query",
"10 runs per planner; metrics: success, time, path length, nodes"],size=18,gap=9)
s=slide("5. Results: CAD and ROS 2 package")
pic(s,R+"ASM00_front.png",0.6,1.5,w=4.0);pic(s,R+"ASM00_side.png",4.7,1.5,w=4.0);pic(s,R+"ASM00_top.png",8.8,1.5,w=4.0)
bullets(s,["19 parts, 2 assemblies, 1 STEP file, 19 STL meshes","Standalone URDF (20 links, 19 joints) and a drop-in xacro for Gazebo / ros2_control","RViz launch with joint sliders; MoveIt 2 and Nav2 configs work unchanged"],y=4.5,size=19)
s=slide("5. Results: arm workspace");pic(s,O+"workspace.png",0.7,1.5,w=8.2)
bullets(s,[f"Max reach from arm base (link1): {f(W['reach_max']*100,1)} cm",f"Max height above floor: {f(W['zmax']*100+1,1)} cm",f"Max forward reach: {f(W['xmax']*100,1)} cm from base_link","20 000 sampled configurations"],x=9.2,y=1.8,w=3.8,size=18)
def simslide(title,fn,pts):
    s=slide(title);pic(s,O+fn,0.4,1.2,w=12.5);bullets(s,pts,y=5.0,size=15,gap=5)
simslide("5. Simulation: RRT","sim_rrt.png",[f"Tree of only {SIM['rrt_nodes']} nodes in this run ({f(rr['time_mean']*1000,0)} ms on average over 10 runs)",f"Path {f(rr['len_mean'])} +/- {f(rr['len_sd'])} rad: {over(rr['len_mean'])} % above the lower bound, jagged",
"Why: it extends towards the nearest node and stops at the first connection, so it is fast but never improves the path"])
simslide("5. Simulation: RRT*","sim_rrtstar.png",[f"About {rs['nodes_mean']:.0f} nodes after 10 s; first path after {f(rs['time_mean']*1000,0)} ms",f"Path {f(rs['len_mean'])} +/- {f(rs['len_sd'])} rad: only {over(rs['len_mean'])} % above the lower bound, smooth",
"Why: best-parent choice and rewiring remove detours; cost is paid in computation (tree about 46x larger than RRT)"])
s=slide("5. RRT* improves with time");pic(s,O+"rrtstar_convergence.png",0.8,1.4,w=7.6)
bullets(s,["Each curve starts at the first path and steps down as rewiring finds shorter routes","The straight line (2.40 rad) is blocked, so the true optimum is above it","Can be stopped at any time and still returns a valid path"],x=8.7,y=1.8,w=4.3,size=17)
simslide("5. Simulation: PRM","sim_prm.png",[f"1500 nodes, {SIM['prm_edges']} edges in this run; {f(pm['time_mean'],1)} s to build, then a graph search",f"Path {f(pm['len_mean'])} +/- {f(pm['len_sd'])} rad: limited to roadmap edges",
"Why: effort is spent up front; later queries in the same static scene are almost free"])
s=slide("5. PRM: effect of roadmap size")
tbl(s,[(str(x["N"]),f"{x['success']}/10",f(x["time"],1),f(x["len"])) for x in prm],["Nodes N","Success","Build + query (s)","Mean path length (rad)"],0.7,1.5,7.2,[1.4,1.4,2.2,2.2],size=15,rh=0.5)
bullets(s,["Build time rises steadily with N","Path length tends to fall as N grows, but neighbouring sizes differ by less than the run-to-run spread (about 0.27 rad)","All sizes solved the problem 10/10"],x=8.3,y=1.5,w=4.7,size=16)
s=slide("5. Results: comparison")
rows=[(k,f"{S[k]['success']}/10",f(S[k]['time_mean']*1000,0),f(S[k]['len_mean']),f"{over(S[k]['len_mean'])} %",f"{S[k]['nodes_mean']:.0f}") for k in ("RRT","RRT*","PRM")]
tbl(s,rows,["Algorithm","Success","First solution (ms)","Path length (rad)","Above lower bound","Nodes"],0.7,1.4,11.9,[2.0,1.5,2.4,2.4,2.2,1.4],size=15,rh=0.45)
pic(s,O+"bench.png",0.7,3.5,w=8.6);bullets(s,["RRT: fastest","RRT*: shortest, most repeatable","PRM: slow build, reusable"],x=9.6,y=3.9,w=3.4,size=17)
tb(s,0.7,6.6,12,0.4,"Our own Python implementation (not OMPL); mean of 10 runs. PRM time = roadmap build + query.",12,False,GREY)
s=slide("5. Justification: which algorithm when?")
tbl(s,[("RRT","One-off motion, need an answer fast","Paths long and jagged; no improvement"),("RRT*","One-off motion, path quality matters; default for this robot","Uses full time budget; larger tree"),("PRM","Repeated moves in a fixed scene (workstation)","Slow first query; rebuild if obstacles move")],["Algorithm","Best used when","Trade-off"],0.7,1.5,11.9,[1.8,5.5,4.6],size=16,rh=0.7)
bullets(s,["Mobile manipulator arrives at new places, so the scene changes: RRT / RRT* fit best","The MoveIt 2 study will test whether RRT* stays the best default"],y=4.6,size=18)
s=slide("5. Limitations");bullets(s,["Simple scene (3 spheres); fixed base; no self-collision, gripper fingers or mesh collision","Our own Python planners, not OMPL: times not comparable with MoveIt","10 runs per planner: rough picture of the spread","Dynamics shown as equations only, not simulated","So: repeat in MoveIt 2 + OMPL + Gazebo"],size=20)
s=slide("6. Conclusion and future work")
bullets(s,["Delivered: parametric CAD + ROS 2 model whose kinematics match the official robot","Simulated and justified RRT, RRT* and PRM on it: RRT quickest, RRT* best paths, PRM reusable","Next: MoveIt 2 + OMPL + Gazebo (RRTkConfigDefault, RRTstarkConfigDefault, PRMkConfigDefault)","Next: RRT-Connect, Informed RRT*, arm dynamics with computed-torque control, base + arm planning, real-robot pick-and-place"],size=20)
s=slide("7. References")
b=s.shapes.add_textbox(Inches(0.7),Inches(1.2),Inches(12),Inches(5.7));tf=b.text_frame;tf.word_wrap=True
for k,r in enumerate(REFS):
    p=tf.paragraphs[0] if k==0 else tf.add_paragraph();ru=p.add_run();ru.text=f"[{k+1}] {r}";ru.font.size=P(10.5);ru.font.color.rgb=INK;p.space_after=P(3)
s=slide("Thank you");tb(s,0.7,2.8,12,1.2,"Thank you  |  Questions?",44,True)
pr.save(O+"Project_Review_Presentation.pptx")
print("built",EQN[0],"equations",n[0],"slides")
